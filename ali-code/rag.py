# ==========================================
#  RAG 模块 v2：支持 TF-IDF 和 pgvector 双模式
#  - 默认 TF-IDF：零配置，开箱即用
#  - pgvector 模式：需要 PostgreSQL + pgvector + embedding API
# ==========================================
import os
import re
import math
from collections import Counter
from abc import ABC, abstractmethod


# ==========================================
#  文档加载器
# ==========================================
def load_pdf(filepath: str) -> str:
    try:
        import PyPDF2
        text = ""
        with open(filepath, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        return text
    except ImportError:
        raise RuntimeError("缺少 PyPDF2 库，请运行: pip install PyPDF2")


def load_word(filepath: str) -> str:
    try:
        from docx import Document
        doc = Document(filepath)
        text = ""
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    text += row_text + "\n"
        return text
    except ImportError:
        raise RuntimeError("缺少 python-docx 库，请运行: pip install python-docx")


def load_txt(filepath: str) -> str:
    encodings = ["utf-8", "gbk", "gb2312", "utf-16"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def load_document(filepath: str) -> str:
    ext = os.path.splitext(filepath)[1].lower()
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"文件不存在: {filepath}")
    if ext == ".pdf":
        return load_pdf(filepath)
    elif ext == ".docx":
        return load_word(filepath)
    elif ext in (".txt", ".md", ".py", ".js", ".json", ".csv"):
        return load_txt(filepath)
    else:
        try:
            return load_txt(filepath)
        except:
            raise ValueError(f"不支持的文件格式: {ext}")


def split_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
    paragraphs = re.split(r'\n\s*\n', text)
    chunks = []
    current_chunk = ""

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(current_chunk) + len(para) + 2 <= chunk_size:
            if current_chunk:
                current_chunk += "\n" + para
            else:
                current_chunk = para
        else:
            if current_chunk:
                chunks.append(current_chunk)
            if len(para) > chunk_size:
                sentences = re.split(r'(?<=[。！？.!?])\s*', para)
                current_chunk = ""
                for sent in sentences:
                    if len(current_chunk) + len(sent) <= chunk_size:
                        current_chunk += sent
                    else:
                        if current_chunk:
                            chunks.append(current_chunk)
                        if len(sent) > chunk_size:
                            for i in range(0, len(sent), chunk_size - overlap):
                                chunks.append(sent[i:i + chunk_size])
                            current_chunk = ""
                        else:
                            current_chunk = sent
            else:
                if overlap > 0 and len(current_chunk) > overlap:
                    current_chunk = current_chunk[-overlap:] + "\n" + para
                else:
                    current_chunk = para

    if current_chunk:
        chunks.append(current_chunk)
    return chunks


# ==========================================
#  抽象基类：检索器接口
# ==========================================
class BaseRetriever(ABC):
    """检索器抽象基类"""

    @abstractmethod
    def add_document(self, filepath: str, chunk_size: int = 500, overlap: int = 50) -> int:
        """添加文档，返回块数"""
        pass

    @abstractmethod
    def add_directory(self, dirpath: str, chunk_size: int = 500, overlap: int = 50) -> tuple[int, int]:
        """添加目录，返回(文档数, 块数)"""
        pass

    @abstractmethod
    def search(self, query: str, top_k: int = 3) -> list[dict]:
        """检索相关文档块"""
        pass

    @abstractmethod
    def list_documents(self) -> list[str]:
        """列出已加载的文档"""
        pass

    @abstractmethod
    def clear(self):
        """清空知识库"""
        pass

    @property
    @abstractmethod
    def chunk_count(self) -> int:
        """总块数"""
        pass


# ==========================================
#  TF-IDF 检索器（默认，零配置）
# ==========================================
class TfidfRetriever(BaseRetriever):
    """基于 TF-IDF 的关键词检索器"""

    def __init__(self):
        self.chunks: list[str] = []
        self.chunk_sources: list[str] = []
        self.doc_names: list[str] = []
        self._idf: dict[str, float] = {}
        self._tfidf_vectors: list[dict] = []

    def add_document(self, filepath: str, chunk_size: int = 500, overlap: int = 50) -> int:
        filename = os.path.basename(filepath)
        if filename in self.doc_names:
            return 0
        text = load_document(filepath)
        chunks = split_text(text, chunk_size, overlap)
        if not chunks:
            return 0
        self.chunks.extend(chunks)
        self.chunk_sources.extend([filename] * len(chunks))
        self.doc_names.append(filename)
        self._rebuild_index()
        return len(chunks)

    def add_directory(self, dirpath: str, chunk_size: int = 500, overlap: int = 50) -> tuple[int, int]:
        if not os.path.isdir(dirpath):
            raise FileNotFoundError(f"目录不存在: {dirpath}")
        supported_exts = {".pdf", ".docx", ".txt", ".md", ".py", ".js", ".json", ".csv"}
        doc_count = 0
        chunk_count = 0
        for root, dirs, files in os.walk(dirpath):
            for file in files:
                ext = os.path.splitext(file)[1].lower()
                if ext in supported_exts:
                    filepath = os.path.join(root, file)
                    try:
                        n = self.add_document(filepath, chunk_size, overlap)
                        if n > 0:
                            doc_count += 1
                            chunk_count += n
                    except Exception as e:
                        print(f"  跳过 {file}: {e}")
        return doc_count, chunk_count

    def _tokenize(self, text: str) -> list[str]:
        tokens = []
        english_words = re.findall(r'[a-zA-Z_][a-zA-Z0-9_]*', text.lower())
        tokens.extend(english_words)
        chinese_chars = re.findall(r'[\u4e00-\u9fff]', text)
        for i in range(len(chinese_chars) - 1):
            tokens.append(chinese_chars[i] + chinese_chars[i + 1])
        numbers = re.findall(r'\d+', text)
        tokens.extend(numbers)
        return tokens

    def _compute_tf(self, tokens: list[str]) -> dict[str, float]:
        if not tokens:
            return {}
        counter = Counter(tokens)
        total = len(tokens)
        return {word: count / total for word, count in counter.items()}

    def _rebuild_index(self):
        if not self.chunks:
            self._idf = {}
            self._tfidf_vectors = []
            return
        tf_list = []
        for chunk in self.chunks:
            tokens = self._tokenize(chunk)
            tf = self._compute_tf(tokens)
            tf_list.append(tf)
        doc_count = len(self.chunks)
        word_doc_freq = Counter()
        for tf in tf_list:
            for word in tf:
                word_doc_freq[word] += 1
        self._idf = {
            word: math.log((doc_count + 1) / (freq + 1)) + 1
            for word, freq in word_doc_freq.items()
        }
        self._tfidf_vectors = []
        for tf in tf_list:
            vector = {word: tf_val * self._idf.get(word, 0) for word, tf_val in tf.items()}
            norm = math.sqrt(sum(v * v for v in vector.values()))
            if norm > 0:
                vector = {k: v / norm for k, v in vector.items()}
            self._tfidf_vectors.append(vector)

    def search(self, query: str, top_k: int = 3) -> list[dict]:
        if not self._tfidf_vectors:
            return []
        query_tokens = self._tokenize(query)
        query_tf = self._compute_tf(query_tokens)
        query_vector = {word: tf_val * self._idf.get(word, 0) for word, tf_val in query_tf.items()}
        norm = math.sqrt(sum(v * v for v in query_vector.values()))
        if norm > 0:
            query_vector = {k: v / norm for k, v in query_vector.items()}
        scores = []
        for i, doc_vector in enumerate(self._tfidf_vectors):
            score = sum(query_vector.get(word, 0) * doc_vector.get(word, 0) for word in query_vector)
            scores.append((i, score))
        scores.sort(key=lambda x: x[1], reverse=True)
        top_results = scores[:top_k]
        results = []
        for idx, score in top_results:
            if score > 0:
                results.append({
                    "chunk": self.chunks[idx],
                    "source": self.chunk_sources[idx],
                    "score": round(score, 4)
                })
        return results

    def list_documents(self) -> list[str]:
        return self.doc_names.copy()

    def clear(self):
        self.chunks = []
        self.chunk_sources = []
        self.doc_names = []
        self._idf = {}
        self._tfidf_vectors = []

    @property
    def chunk_count(self) -> int:
        return len(self.chunks)


# ==========================================
#  pgvector 检索器（高级模式，需要 PostgreSQL）
# ==========================================
class PgVectorRetriever(BaseRetriever):
    """
    基于 PostgreSQL + pgvector 的向量检索器
    需要：
    1. PostgreSQL 数据库 + pgvector 扩展
    2. Embedding API（豆包的 embedding 接口）
    """

    def __init__(self, db_config: dict, embedding_config: dict):
        """
        Args:
            db_config: {host, port, user, password, dbname}
            embedding_config: {client, model, dimensions}
        """
        self.db_config = db_config
        self.embedding_config = embedding_config
        self._conn = None
        self._init_db()

    def _get_conn(self):
        import psycopg2
        if self._conn is None or self._conn.closed:
            self._conn = psycopg2.connect(**self.db_config)
            self._conn.autocommit = True
        return self._conn

    def _init_db(self):
        """初始化数据库表和扩展"""
        conn = self._get_conn()
        with conn.cursor() as cur:
            # 创建 pgvector 扩展
            cur.execute("CREATE EXTENSION IF NOT EXISTS vector")
            # 创建文档表
            dims = self.embedding_config.get("dimensions", 1536)
            cur.execute(f"""
                CREATE TABLE IF NOT EXISTS document_chunks (
                    id SERIAL PRIMARY KEY,
                    content TEXT NOT NULL,
                    source VARCHAR(500) NOT NULL,
                    doc_name VARCHAR(200) NOT NULL,
                    embedding vector({dims})
                )
            """)
            # 创建索引
            cur.execute("""
                CREATE INDEX IF NOT EXISTS idx_chunks_embedding
                ON document_chunks
                USING hnsw (embedding vector_cosine_ops)
            """)

    def _get_embedding(self, text: str) -> list[float]:
        """调用 embedding API 获取向量"""
        client = self.embedding_config["client"]
        model = self.embedding_config["model"]
        response = client.embeddings.create(
            model=model,
            input=text
        )
        return response.data[0].embedding

    def add_document(self, filepath: str, chunk_size: int = 500, overlap: int = 50) -> int:
        filename = os.path.basename(filepath)

        # 检查是否已加载
        conn = self._get_conn()
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) FROM document_chunks WHERE doc_name = %s", (filename,))
            if cur.fetchone()[0] > 0:
                return 0

        text = load_document(filepath)
        chunks = split_text(text, chunk_size, overlap)
        if not chunks:
            return 0

        # 批量生成 embedding 并存入数据库
        conn = self._get_conn()
        with conn.cursor() as cur:
            for chunk in chunks:
                embedding = self._get_embedding(chunk)
                embedding_str = "[" + ",".join(str(x) for x in embedding) + "]"
                cur.execute(
                    "INSERT INTO document_chunks (content, source, doc_name, embedding) VALUES (%s, %s, %s, %s::vector)",
                    (chunk, filename, filename, embedding_str)
                )

        return len(chunks)

    def add_directory(self, dirpath: str, chunk_size: int = 500, overlap: int = 50) -> tuple[int, int]:
        if not os.path.isdir(dirpath):
            raise FileNotFoundError(f"目录不存在: {dirpath}")
        supported_exts = {".pdf", ".docx", ".txt", ".md", ".py", ".js", ".json", ".csv"}
        doc_count = 0
        chunk_count = 0
        for root, dirs, files in os.walk(dirpath):
            for file in files:
                ext = os.path.splitext(file)[1].lower()
                if ext in supported_exts:
                    filepath = os.path.join(root, file)
                    try:
                        n = self.add_document(filepath, chunk_size, overlap)
                        if n > 0:
                            doc_count += 1
                            chunk_count += n
                    except Exception as e:
                        print(f"  跳过 {file}: {e}")
        return doc_count, chunk_count

    def search(self, query: str, top_k: int = 3) -> list[dict]:
        query_embedding = self._get_embedding(query)
        embedding_str = "[" + ",".join(str(x) for x in query_embedding) + "]"

        conn = self._get_conn()
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT content, source, 1 - (embedding <=> %s::vector) as score
                FROM document_chunks
                ORDER BY embedding <=> %s::vector
                LIMIT %s
                """,
                (embedding_str, embedding_str, top_k)
            )
            rows = cur.fetchall()

        results = []
        for content, source, score in rows:
            results.append({
                "chunk": content,
                "source": source,
                "score": round(score, 4)
            })
        return results

    def list_documents(self) -> list[str]:
        conn = self._get_conn()
        with conn.cursor() as cur:
            cur.execute("SELECT DISTINCT doc_name FROM document_chunks ORDER BY doc_name")
            rows = cur.fetchall()
        return [row[0] for row in rows]

    def clear(self):
        conn = self._get_conn()
        with conn.cursor() as cur:
            cur.execute("TRUNCATE TABLE document_chunks")

    @property
    def chunk_count(self) -> int:
        conn = self._get_conn()
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) FROM document_chunks")
            return cur.fetchone()[0]


# ==========================================
#  全局检索器工厂
# ==========================================
_retriever_instance = None


def init_retriever(mode: str = "tfidf", **kwargs) -> BaseRetriever:
    """
    初始化全局检索器
    Args:
        mode: "tfidf" 或 "pgvector"
        **kwargs: 对应检索器的配置参数
    """
    global _retriever_instance
    if mode == "tfidf":
        _retriever_instance = TfidfRetriever()
    elif mode == "pgvector":
        _retriever_instance = PgVectorRetriever(**kwargs)
    else:
        raise ValueError(f"未知的检索模式: {mode}")
    return _retriever_instance


def get_retriever() -> BaseRetriever:
    """获取全局检索器实例（如果未初始化，默认用 TF-IDF）"""
    global _retriever_instance
    if _retriever_instance is None:
        _retriever_instance = TfidfRetriever()
    return _retriever_instance
